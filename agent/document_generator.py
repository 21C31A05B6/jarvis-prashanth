"""Generates downloadable documents from plain text: .docx (python-docx)
and .pdf (PyMuPDF, already a project dependency used elsewhere for PDF
handling). PowerPoint generation already exists separately and more richly
via actions/presentation_maker.py (desktop-only); Excel export is a natural
follow-up, not yet built.
"""

from __future__ import annotations

import io


def generate_docx(title: str, content: str) -> bytes:
    from docx import Document

    document = Document()
    document.add_heading(title, level=1)
    for paragraph in content.split("\n\n"):
        paragraph = paragraph.strip()
        if not paragraph:
            continue
        if paragraph.startswith("- ") or paragraph.startswith("* "):
            for line in paragraph.splitlines():
                line = line.strip().lstrip("-*").strip()
                if line:
                    document.add_paragraph(line, style="List Bullet")
        else:
            document.add_paragraph(paragraph)

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def generate_pdf(title: str, content: str) -> bytes:
    import pymupdf

    document = pymupdf.open()
    margin = 50
    fontsize = 11

    def new_page():
        page = document.new_page()
        rect = pymupdf.Rect(margin, margin, page.rect.width - margin, page.rect.height - margin)
        return page, rect

    def fits(page, rect, text: str) -> bool:
        # PyMuPDF's render_mode=3 ("invisible") still writes an extractable
        # text object into the page — it's only invisible when displayed,
        # not absent from the content stream. Measuring against it would
        # leave duplicate invisible text baked into the real output. Measure
        # against a fully disposable scratch document instead.
        scratch = pymupdf.open()
        scratch_page = scratch.new_page(width=page.rect.width, height=page.rect.height)
        overflow = scratch_page.insert_textbox(rect, text, fontsize=fontsize, fontname="helv")
        scratch.close()
        return overflow >= 0

    page, rect = new_page()
    accumulated = ""
    paragraphs = [title, "", *content.split("\n\n")]

    for paragraph in paragraphs:
        candidate = f"{accumulated}\n\n{paragraph}" if accumulated else paragraph
        if fits(page, rect, candidate):
            accumulated = candidate
            continue

        if accumulated:
            page.insert_textbox(rect, accumulated, fontsize=fontsize, fontname="helv")
        page, rect = new_page()
        accumulated = paragraph

        # A single paragraph too large for one page: binary-search the
        # longest word-aligned prefix that fits, spilling the rest forward.
        while not fits(page, rect, accumulated):
            words = accumulated.split(" ")
            if len(words) <= 1:
                break
            lo, hi = 1, len(words)
            while lo < hi:
                mid = (lo + hi + 1) // 2
                if fits(page, rect, " ".join(words[:mid])):
                    lo = mid
                else:
                    hi = mid - 1
            page.insert_textbox(rect, " ".join(words[:lo]), fontsize=fontsize, fontname="helv")
            page, rect = new_page()
            accumulated = " ".join(words[lo:])

    if accumulated:
        page.insert_textbox(rect, accumulated, fontsize=fontsize, fontname="helv")

    buffer = io.BytesIO()
    document.save(buffer)
    document.close()
    return buffer.getvalue()
